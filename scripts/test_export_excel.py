"""
Automated Test Suite for Export to Excel Feature
Validates:
1. ExportService generates styled XLSX with correct headers and total rows.
2. ExportService generates valid CSV with UTF-8 BOM and total rows.
3. Proper handling of year, month, and transaction_type filters.
4. HTML contains modal-export-report with all required filter controls.
5. app.js binds export buttons and prevents browser page reloads.
6. api.js contains exportTransactions method.
"""
import io
import os
import sys
import re

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.services.export_service import ExportService


def run_tests():
    print("=" * 65)
    print("TESTING EXPORT TO EXCEL & CSV FEATURE")
    print("=" * 65)

    sample_transactions = [
        {
            "id": 101,
            "type": "EXPENSE",
            "amount": 35000,
            "description": "Makan Siang Soto Ayam",
            "transaction_date": "2026-10-02",
            "category_name": "Makanan & Minuman",
            "source_account_name": "BCA",
            "destination_account_name": None,
        },
        {
            "id": 102,
            "type": "EXPENSE",
            "amount": 120000,
            "description": "Belanja Mingguan Indomaret",
            "transaction_date": "2026-10-05",
            "category_name": "Kebutuhan Pokok",
            "source_account_name": "BCA",
            "destination_account_name": None,
        },
        {
            "id": 103,
            "type": "INCOME",
            "amount": 10000000,
            "description": "Gaji Bulanan",
            "transaction_date": "2026-10-01",
            "category_name": "Gaji",
            "source_account_name": None,
            "destination_account_name": "BCA",
        },
    ]

    # Test 1: XLSX Generation with ALL filter
    data_all, fn_all, mt_all = ExportService.export(
        sample_transactions, year=2026, month=10, transaction_type="ALL", format_type="xlsx"
    )
    assert len(data_all) > 1000, "XLSX output must be a valid binary file"
    assert fn_all.endswith(".xlsx"), f"Filename must end with .xlsx, got: {fn_all}"
    assert "application/vnd.openxmlformats" in mt_all
    print(f" [PASS] 1. XLSX (All Transactions) generated: {fn_all} ({len(data_all)} bytes)")

    # Test 2: XLSX Generation with EXPENSE filter (only expenses)
    expense_txs = [t for t in sample_transactions if t["type"] == "EXPENSE"]
    data_exp, fn_exp, mt_exp = ExportService.export(
        expense_txs, year=2026, month=10, transaction_type="EXPENSE", format_type="xlsx"
    )
    assert "Pengeluaran" in fn_exp, f"Filename must mention Pengeluaran, got: {fn_exp}"
    assert len(data_exp) > 1000
    print(f" [PASS] 2. XLSX (Khusus Pengeluaran) generated: {fn_exp} ({len(data_exp)} bytes)")

    # Test 3: Inspect XLSX content structure using openpyxl
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data_exp))
    ws = wb.active
    assert ws.title == "Laporan Transaksi"

    # Header check
    headers = [cell.value for cell in ws[5] if cell.value is not None]
    assert "Tanggal Transaksi" in headers, "Tanggal Transaksi missing in headers"
    assert "Kategori" in headers, "Kategori missing in headers"
    assert "Deskripsi / Keterangan" in headers, "Deskripsi missing in headers"
    assert "Nominal Pengeluaran (Rp)" in headers, "Nominal Pengeluaran missing in headers"

    # Total row check
    found_total = False
    for r in range(6, ws.max_row + 1):
        val = str(ws.cell(row=r, column=1).value or "")
        if "TOTAL" in val:
            found_total = True
            # Check sum formula in last column
            formula_val = str(ws.cell(row=r, column=6).value or "")
            assert formula_val.startswith("=SUM"), f"Total cell must contain formula, got: {formula_val}"
            break
    assert found_total, "TOTAL row must exist at the bottom of the table!"
    print(" [PASS] 3. XLSX table columns and bottom TOTAL formula verified")

    # Test 4: CSV Generation with UTF-8 BOM and Total row
    data_csv, fn_csv, mt_csv = ExportService.export(
        expense_txs, year=2026, month=10, transaction_type="EXPENSE", format_type="csv"
    )
    assert data_csv.startswith(b"\xef\xbb\xbf"), "CSV must start with UTF-8 BOM (\\ufeff) for Excel compatibility"
    csv_str = data_csv.decode("utf-8")
    assert "TOTAL PENGELUARAN" in csv_str, "TOTAL PENGELUARAN row must exist in CSV"
    assert "155000" in csv_str, "Total sum (35000 + 120000 = 155000) must appear in CSV"
    print(f" [PASS] 4. CSV generation with UTF-8 BOM and bottom Total verified: {fn_csv}")

    # Test 5: Verify modal-export-report in index.html
    with open("app/static/index.html", "r", encoding="utf-8") as f:
        html = f.read()

    assert 'id="modal-export-report"' in html, "modal-export-report missing in index.html"
    assert 'id="form-export-report"' in html, "form-export-report missing in index.html"
    assert 'id="select-export-type"' in html, "select-export-type missing in index.html"
    assert 'id="select-export-year"' in html, "select-export-year missing in index.html"
    assert 'id="select-export-month"' in html, "select-export-month missing in index.html"
    assert 'id="select-export-format"' in html, "select-export-format missing in index.html"
    assert 'id="btn-cancel-export"' in html, "btn-cancel-export missing in index.html"
    assert 'id="btn-confirm-export"' in html, "btn-confirm-export missing in index.html"
    print(" [PASS] 5. modal-export-report DOM elements and filter controls verified in index.html")

    # Test 6: Verify app.js bindings and handler methods
    with open("app/static/js/app.js", "r", encoding="utf-8") as f:
        js = f.read()

    assert "this.modalExportReport = document.getElementById('modal-export-report');" in js
    assert "this.openExportModal();" in js
    assert "async handleExportSubmit(e)" in js
    assert "targetId === 'form-export-report'" in js, "Delegated submit preventDefault for export form missing"
    print(" [PASS] 6. app.js initialization, openExportModal, and delegated submit handlers verified")

    # Test 7: Verify api.js exportTransactions client method
    with open("app/static/js/api.js", "r", encoding="utf-8") as f:
        api_js = f.read()

    assert "exportTransactions: async (params = {}) =>" in api_js
    assert "/transactions/export" in api_js
    print(" [PASS] 7. api.js exportTransactions endpoint client verified")

    # Test 8: Verify backend router in app/api/transactions.py
    with open("app/api/transactions.py", "r", encoding="utf-8") as f:
        tx_router = f.read()

    assert '@router.get("/export")' in tx_router, "/export route missing in transactions.py"
    assert "TransactionService.get_transactions_for_export" in tx_router
    assert "ExportService.export" in tx_router
    print(" [PASS] 8. Backend /transactions/export route and service integration verified")

    print("=" * 65)
    print("ALL 8 EXPORT FEATURE TESTS PASSED (100% GREEN)")
    print("=" * 65)


if __name__ == "__main__":
    run_tests()
