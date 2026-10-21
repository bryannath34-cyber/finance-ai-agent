import io
import csv
from datetime import datetime, date
from decimal import Decimal
from typing import List, Dict, Any, Optional, Tuple

MONTH_NAMES = {
    1: "Januari", 2: "Februari", 3: "Maret", 4: "April",
    5: "Mei", 6: "Juni", 7: "Juli", 8: "Agustus",
    9: "September", 10: "Oktober", 11: "November", 12: "Desember"
}

class ExportService:
    @staticmethod
    def _format_date(val: Any) -> str:
        if isinstance(val, (datetime, date)):
            return val.strftime("%Y-%m-%d")
        return str(val or "")

    @staticmethod
    def _build_meta_strings(
        year: Optional[int] = None,
        month: Optional[int] = None,
        transaction_type: Optional[str] = None,
        format_ext: str = "xlsx"
    ) -> Tuple[str, str, str]:
        """Menghasilkan teks periode, teks filter, dan nama berkas yang rapi."""
        # Periode
        if month and year:
            month_name = MONTH_NAMES.get(month, f"Bulan-{month}")
            period_str = f"{month_name} {year}"
            period_slug = f"{month_name}_{year}"
        elif year:
            period_str = f"Tahun {year}"
            period_slug = f"Tahun_{year}"
        else:
            period_str = "Semua Periode"
            period_slug = "Semua_Periode"

        # Tipe Transaksi
        tx_type = (transaction_type or "ALL").upper()
        if tx_type == "EXPENSE":
            filter_str = "Khusus Pengeluaran"
            type_slug = "Pengeluaran"
        elif tx_type == "INCOME":
            filter_str = "Khusus Pemasukan"
            type_slug = "Pemasukan"
        elif tx_type == "TRANSFER":
            filter_str = "Khusus Transfer"
            type_slug = "Transfer"
        else:
            filter_str = "Semua Transaksi"
            type_slug = "Transaksi"

        filename = f"Laporan_{type_slug}_{period_slug}.{format_ext}"
        return period_str, filter_str, filename

    @classmethod
    def generate_excel(
        cls,
        transactions: List[Dict[str, Any]],
        year: Optional[int] = None,
        month: Optional[int] = None,
        transaction_type: Optional[str] = None
    ) -> Tuple[bytes, str]:
        """
        Menghasilkan berkas Excel (.xlsx) rapi dengan header bertema Nordic Green,
        kolom tanggal, kategori, rekening, deskripsi, nominal (pemasukan/pengeluaran),
        dan baris TOTAL di bagian paling bawah tabel.
        """
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        period_str, filter_str, filename = cls._build_meta_strings(year, month, transaction_type, "xlsx")
        tx_type = (transaction_type or "ALL").upper()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Laporan Transaksi"
        ws.views.sheetView[0].showGridLines = True

        # Styles definition
        font_name = "Segoe UI"
        primary_dark = "1B3B2B"  # Nordic Dark Green
        subtle_green = "EAF2EC"  # Light sage green for total
        zebra_light = "F8FAF8"
        border_color = "D0D7DE"

        thin_side = Side(border_style="thin", color=border_color)
        double_side = Side(border_style="double", color=primary_dark)
        top_total_side = Side(border_style="thin", color=primary_dark)

        border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        border_total = Border(left=thin_side, right=thin_side, top=top_total_side, bottom=double_side)

        header_fill = PatternFill(start_color=primary_dark, end_color=primary_dark, fill_type="solid")
        header_font = Font(name=font_name, size=11, bold=True, color="FFFFFF")
        total_fill = PatternFill(start_color=subtle_green, end_color=subtle_green, fill_type="solid")

        # Title Banner
        ws["A1"] = "FINANCE AI AGENT — LAPORAN TRANSAKSI"
        ws["A1"].font = Font(name=font_name, size=14, bold=True, color=primary_dark)

        ws["A2"] = f"Periode: {period_str}  |  Filter: {filter_str}  |  Total Data: {len(transactions)} transaksi"
        ws["A2"].font = Font(name=font_name, size=10, italic=True, color="555555")

        now_str = datetime.now().strftime("%d-%m-%Y %H:%M:%S")
        ws["A3"] = f"Dicetak pada: {now_str}"
        ws["A3"].font = Font(name=font_name, size=9, color="777777")

        # Table Headers Configuration
        is_only_expense = (tx_type == "EXPENSE")
        is_only_income = (tx_type == "INCOME")

        if is_only_expense:
            headers = [
                "No", "Tanggal Transaksi", "Kategori", "Rekening",
                "Deskripsi / Keterangan", "Nominal Pengeluaran (Rp)"
            ]
        elif is_only_income:
            headers = [
                "No", "Tanggal Transaksi", "Kategori", "Rekening",
                "Deskripsi / Keterangan", "Nominal Pemasukan (Rp)"
            ]
        else:
            headers = [
                "No", "Tanggal Transaksi", "Tipe Transaksi", "Kategori",
                "Rekening", "Deskripsi / Keterangan", "Pemasukan (Rp)", "Pengeluaran (Rp)"
            ]

        header_row_idx = 5
        ws.row_dimensions[header_row_idx].height = 26

        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=header_row_idx, column=col_idx, value=h_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_cell

        # Populate Data Rows
        start_data_row = 6
        current_row = start_data_row

        sum_income = Decimal("0.00")
        sum_expense = Decimal("0.00")
        sum_single = Decimal("0.00")

        for idx, tx in enumerate(transactions, start=1):
            t_date = cls._format_date(tx.get("transaction_date"))
            raw_type = (tx.get("type") or "").upper()
            t_type = "Pengeluaran" if raw_type == "EXPENSE" else ("Pemasukan" if raw_type == "INCOME" else "Transfer")
            cat_name = tx.get("category_name") or "-"
            
            # Rekening
            src_acc = tx.get("source_account_name")
            dst_acc = tx.get("destination_account_name")
            if raw_type == "EXPENSE":
                acc_name = src_acc or "-"
            elif raw_type == "INCOME":
                acc_name = dst_acc or "-"
            else:
                acc_name = f"{src_acc or '-'} -> {dst_acc or '-'}" if (src_acc or dst_acc) else "-"

            desc = tx.get("description") or "-"
            raw_amt = tx.get("amount") or 0
            try:
                amt = Decimal(str(raw_amt))
            except Exception:
                amt = Decimal("0.00")

            ws.row_dimensions[current_row].height = 20
            row_fill = PatternFill(start_color=zebra_light, end_color=zebra_light, fill_type="solid") if (idx % 2 == 0) else None

            if is_only_expense:
                row_vals = [idx, t_date, cat_name, acc_name, desc, float(amt)]
                sum_single += amt
            elif is_only_income:
                row_vals = [idx, t_date, cat_name, acc_name, desc, float(amt)]
                sum_single += amt
            else:
                inc_val = float(amt) if raw_type == "INCOME" else 0.0
                exp_val = float(amt) if raw_type == "EXPENSE" else 0.0
                if raw_type == "INCOME":
                    sum_income += amt
                elif raw_type == "EXPENSE":
                    sum_expense += amt
                row_vals = [idx, t_date, t_type, cat_name, acc_name, desc, inc_val, exp_val]

            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = Font(name=font_name, size=10, color="222222")
                cell.border = border_cell
                if row_fill:
                    cell.fill = row_fill

                # Alignment & number formats
                if col_idx == 1:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 2 or (not is_only_expense and not is_only_income and col_idx == 3):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif (is_only_expense or is_only_income) and col_idx == 6:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0"
                elif not is_only_expense and not is_only_income and col_idx in (7, 8):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0"
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            current_row += 1

        # Handle empty dataset
        if len(transactions) == 0:
            ws.row_dimensions[current_row].height = 24
            cell = ws.cell(row=current_row, column=1, value="Tidak ada transaksi pada periode yang dipilih.")
            cell.font = Font(name=font_name, size=10, italic=True, color="777777")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
            for c in range(1, len(headers) + 1):
                ws.cell(row=current_row, column=c).border = border_cell
            current_row += 1

        # Bottom Total Row
        ws.row_dimensions[current_row].height = 24
        last_data_row = max(start_data_row, current_row - 1)

        if is_only_expense or is_only_income:
            total_label = "TOTAL PENGELUARAN" if is_only_expense else "TOTAL PEMASUKAN"
            ws.cell(row=current_row, column=1, value=total_label)
            num_col_letter = get_column_letter(6)
            ws.cell(row=current_row, column=6, value=f"=SUM({num_col_letter}{start_data_row}:{num_col_letter}{last_data_row})")
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)

            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=c)
                cell.fill = total_fill
                cell.font = Font(name=font_name, size=11, bold=True, color=primary_dark)
                cell.border = border_total
                if c == 1:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c == 6:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0"
        else:
            ws.cell(row=current_row, column=1, value="TOTAL")
            col_inc_letter = get_column_letter(7)
            col_exp_letter = get_column_letter(8)
            ws.cell(row=current_row, column=7, value=f"=SUM({col_inc_letter}{start_data_row}:{col_inc_letter}{last_data_row})")
            ws.cell(row=current_row, column=8, value=f"=SUM({col_exp_letter}{start_data_row}:{col_exp_letter}{last_data_row})")
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)

            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=c)
                cell.fill = total_fill
                cell.font = Font(name=font_name, size=11, bold=True, color=primary_dark)
                cell.border = border_total
                if c == 1:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c in (7, 8):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0"

            # Row Arus Kas Bersih
            current_row += 1
            ws.row_dimensions[current_row].height = 24
            ws.cell(row=current_row, column=1, value="ARUS KAS BERSIH (NET)")
            ws.cell(row=current_row, column=7, value=f"={col_inc_letter}{current_row - 1}-{col_exp_letter}{current_row - 1}")
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
            ws.merge_cells(start_row=current_row, start_column=7, end_row=current_row, end_column=8)

            net_fill = PatternFill(start_color="DCE8E0", end_color="DCE8E0", fill_type="solid")
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=current_row, column=c)
                cell.fill = net_fill
                cell.font = Font(name=font_name, size=11, bold=True, color=primary_dark)
                cell.border = border_total
                if c == 1:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c in (7, 8):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0"

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            # Skip title rows 1..3 for calculating column widths
            for cell in col[4:]:
                val_str = str(cell.value or "")
                if val_str.startswith("="):
                    val_str = "Rp 999.999.999"
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        ws.column_dimensions["A"].width = 7  # No
        ws.column_dimensions["B"].width = 16 # Tanggal

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue(), filename

    @classmethod
    def generate_csv(
        cls,
        transactions: List[Dict[str, Any]],
        year: Optional[int] = None,
        month: Optional[int] = None,
        transaction_type: Optional[str] = None
    ) -> Tuple[bytes, str]:
        """
        Menghasilkan berkas CSV rapi dengan UTF-8 BOM, judul laporan, kolom rapi,
        serta baris TOTAL di bagian paling bawah.
        """
        period_str, filter_str, filename = cls._build_meta_strings(year, month, transaction_type, "csv")
        tx_type = (transaction_type or "ALL").upper()
        is_only_expense = (tx_type == "EXPENSE")
        is_only_income = (tx_type == "INCOME")

        output = io.StringIO()
        writer = csv.writer(output, dialect="excel")

        # Meta info
        writer.writerow(["FINANCE AI AGENT - LAPORAN TRANSAKSI"])
        writer.writerow([f"Periode: {period_str}", f"Filter: {filter_str}", f"Total: {len(transactions)} transaksi"])
        writer.writerow([f"Dicetak pada: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
        writer.writerow([])

        if is_only_expense:
            headers = ["No", "Tanggal Transaksi", "Kategori", "Rekening", "Deskripsi", "Nominal Pengeluaran (Rp)"]
        elif is_only_income:
            headers = ["No", "Tanggal Transaksi", "Kategori", "Rekening", "Deskripsi", "Nominal Pemasukan (Rp)"]
        else:
            headers = ["No", "Tanggal Transaksi", "Tipe", "Kategori", "Rekening", "Deskripsi", "Pemasukan (Rp)", "Pengeluaran (Rp)"]

        writer.writerow(headers)

        sum_income = Decimal("0.00")
        sum_expense = Decimal("0.00")
        sum_single = Decimal("0.00")

        for idx, tx in enumerate(transactions, start=1):
            t_date = cls._format_date(tx.get("transaction_date"))
            raw_type = (tx.get("type") or "").upper()
            t_type = "Pengeluaran" if raw_type == "EXPENSE" else ("Pemasukan" if raw_type == "INCOME" else "Transfer")
            cat_name = tx.get("category_name") or "-"

            src_acc = tx.get("source_account_name")
            dst_acc = tx.get("destination_account_name")
            if raw_type == "EXPENSE":
                acc_name = src_acc or "-"
            elif raw_type == "INCOME":
                acc_name = dst_acc or "-"
            else:
                acc_name = f"{src_acc or '-'} -> {dst_acc or '-'}" if (src_acc or dst_acc) else "-"

            desc = tx.get("description") or "-"
            raw_amt = tx.get("amount") or 0
            try:
                amt = Decimal(str(raw_amt))
            except Exception:
                amt = Decimal("0.00")

            if is_only_expense or is_only_income:
                sum_single += amt
                writer.writerow([idx, t_date, cat_name, acc_name, desc, int(amt)])
            else:
                inc_val = int(amt) if raw_type == "INCOME" else 0
                exp_val = int(amt) if raw_type == "EXPENSE" else 0
                if raw_type == "INCOME":
                    sum_income += amt
                elif raw_type == "EXPENSE":
                    sum_expense += amt
                writer.writerow([idx, t_date, t_type, cat_name, acc_name, desc, inc_val, exp_val])

        # Bottom Total
        writer.writerow([])
        if is_only_expense or is_only_income:
            lbl = "TOTAL PENGELUARAN" if is_only_expense else "TOTAL PEMASUKAN"
            writer.writerow(["", "", "", "", lbl, int(sum_single)])
        else:
            writer.writerow(["", "", "", "", "", "TOTAL", int(sum_income), int(sum_expense)])
            writer.writerow(["", "", "", "", "", "ARUS KAS BERSIH (NET)", int(sum_income - sum_expense)])

        # Prepend UTF-8 BOM so Excel opens with proper encoding
        csv_bytes = ("\ufeff" + output.getvalue()).encode("utf-8")
        return csv_bytes, filename

    @classmethod
    def export(
        cls,
        transactions: List[Dict[str, Any]],
        year: Optional[int] = None,
        month: Optional[int] = None,
        transaction_type: Optional[str] = None,
        format_type: str = "xlsx"
    ) -> Tuple[bytes, str, str]:
        """
        Orkestrasi ekspor: mengembalikan (bytes_data, filename, media_type).
        Secara otomatis beralih ke CSV jika openpyxl belum terinstal.
        """
        fmt = (format_type or "xlsx").lower()
        if fmt == "xlsx":
            try:
                data, filename = cls.generate_excel(transactions, year, month, transaction_type)
                media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                return data, filename, media_type
            except ImportError:
                # Fallback ke CSV jika openpyxl belum ada
                data, filename = cls.generate_csv(transactions, year, month, transaction_type)
                media_type = "text/csv; charset=utf-8"
                return data, filename, media_type

        data, filename = cls.generate_csv(transactions, year, month, transaction_type)
        media_type = "text/csv; charset=utf-8"
        return data, filename, media_type
